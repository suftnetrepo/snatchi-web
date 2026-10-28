from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parents[1] / "docs" / "Snatchi-Product-Overview-for-PlasmaPro.docx"

INDIGO = "4338CA"
DEEP = "242057"
INK = "1D1B35"
MUTED = "66637A"
PALE = "EEECFF"
LIGHT = "F7F6FC"
LINE = "DEDCEA"
WHITE = "FFFFFF"
AMBER = "9B5C08"
RED = "A3342C"
GREEN = "166534"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            cell_margins(cell)


def set_font(run, size=None, bold=None, color=None, name="Aptos"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_font(run, 8.5, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def border_bottom(paragraph, color=INDIGO, size="18"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def add_field(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"{label}  ")
    set_font(r, 9.5, True, INDIGO)
    r = p.add_run(value)
    set_font(r, 9.5, False, INK)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_font(r, 10.5, True, INK)
        r = p.add_run(text[len(bold_lead):])
        set_font(r, 10.5, False, INK)
    else:
        r = p.add_run(text)
        set_font(r, 10.5, False, INK)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(.42)
        p.paragraph_format.first_line_indent = Inches(-.18)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(item)
        set_font(r, 10.25, False, INK)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(.42)
        p.paragraph_format.first_line_indent = Inches(-.18)
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(item)
        set_font(r, 10.25, False, INK)


def add_callout(doc, label, text, tone="info"):
    fills = {"info": PALE, "positive": "EAF7EF", "caution": "FFF5E6", "risk": "FFF0EF"}
    colors = {"info": INDIGO, "positive": GREEN, "caution": AMBER, "risk": RED}
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(.12)
    p.paragraph_format.right_indent = Inches(.12)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(9)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fills[tone])
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "22")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), colors[tone])
    borders.append(left)
    p_pr.append(borders)
    r = p.add_run(label.upper())
    set_font(r, 8.5, True, colors[tone])
    r.add_break()
    r = p.add_run(text)
    set_font(r, 10, False, INK)


def add_table(doc, headers, rows, widths, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, value in enumerate(headers):
        cell = header.cells[idx]
        shade(cell, DEEP)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(value)
        set_font(r, font_size, True, WHITE)
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cell = cells[idx]
            if row_idx % 2:
                shade(cell, LIGHT)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            set_font(r, font_size, False, INK)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_chapter(doc, number, title, deck=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(15)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"{number:02d}  /  PRODUCT OVERVIEW")
    set_font(r, 8.5, True, INDIGO)
    h = doc.add_heading(title, level=1)
    h.paragraph_format.keep_with_next = True
    if deck:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(12)
        r = p.add_run(deck)
        set_font(r, 11, False, MUTED)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(.82)
    section.bottom_margin = Inches(.82)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(.42)
    section.footer_distance = Inches(.42)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, before, after, color in (
        ("Heading 1", 20, 17, 8, DEEP),
        ("Heading 2", 14, 13, 6, INDIGO),
        ("Heading 3", 11.5, 9, 4, DEEP),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = header.add_run("SNATCHI  /  PRODUCT OVERVIEW")
    set_font(r, 8, True, MUTED)
    footer = section.footer.paragraphs[0]
    page_number(footer)

    # Cover — editorial report pattern with restrained brand treatment.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(58)
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("SNATCHI")
    set_font(r, 11, True, INDIGO)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(9)
    r = title.add_run("Product Overview\nand Delivery Position")
    set_font(r, 32, True, DEEP, "Aptos Display")
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(20)
    r = sub.add_run("A connected platform for AV projects, engineering teams, bookings, communication, payments and field delivery.")
    set_font(r, 15, False, MUTED)
    border_bottom(sub, INDIGO, "20")
    add_field(doc, "PREPARED FOR", "PlasmaPro Ltd")
    add_field(doc, "PRODUCT", "Snatchi")
    add_field(doc, "DEVELOPMENT PARTNER", "Suftnet")
    add_field(doc, "DATE", "22 August 2026")
    add_field(doc, "DOCUMENT PURPOSE", "Stakeholder presentation and product-readiness overview")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(85)
    r = p.add_run("CONFIDENTIAL — FOR PLASMAPRO STAKEHOLDER REVIEW")
    set_font(r, 8.5, True, MUTED)
    doc.add_page_break()

    add_chapter(doc, 1, "Executive summary", "What has been built, who it serves and how the product fits together.")
    add_body(doc, "Snatchi is a two-surface operational platform for audiovisual project delivery. The web application gives AV organisations a controlled workspace for projects, people, scheduling, communication, invoices, subscriptions and platform operations. The React Native mobile application gives engineers a field-focused experience for bookings, project information, documents, chat, service rates, invoicing and location-aware job progression.")
    add_body(doc, "The product is owned and operated by PlasmaPro Ltd and developed by Suftnet. Its design connects office coordination with field execution: an organisation creates a project and books an engineer on the web; the engineer responds and performs the work through the mobile app; both parties share status, communication, documentation and financial records through common services.")
    add_callout(doc, "Current delivery position", "The core product and its principal commercial controls are implemented. The web production build passes, plan limits are enforced server-side, Stripe live price identifiers are preserved, and the first admin control-plane release is available. Final production release still requires environment verification, mobile release validation, live-service smoke testing and the explicitly deferred admin MFA/session-security milestone.", "positive")
    doc.add_heading("The value proposition", level=2)
    add_bullets(doc, [
        "One operational record from project creation through booking, field delivery and invoice review.",
        "Clear role separation for platform administrators, organisation owners, managers, engineers and guests.",
        "Internal-team scheduling plus an external engineer network with service rates and marketplace payments.",
        "Web and mobile workflows connected by shared APIs, notifications and booking-status rules.",
        "A realistic three-plan commercial model with enforceable project, member and document allowances.",
        "A safer administrative control plane focused on visibility, reversible access control, payment investigation and auditability."
    ])
    doc.add_heading("Product at a glance", level=2)
    add_table(doc, ["Surface", "Primary audience", "Principal outcome"], [
        ["Public website", "Prospective customers", "Understand the product, compare plans, contact PlasmaPro and start subscription onboarding."],
        ["Organisation web workspace", "Owners and managers", "Plan projects, manage people, book engineers, communicate, review invoices and control billing."],
        ["Engineer mobile app", "Field engineers", "Receive and progress bookings, access job information, communicate, manage documents, rates and invoices."],
        ["Platform administration", "Authorised PlasmaPro operators", "Observe the platform, control organisation access, investigate payment issues and maintain an audit trail."],
    ], [1800, 2300, 5260])

    add_chapter(doc, 2, "Users, roles and operating model", "Each identity has a specific job and a bounded permission set.")
    add_table(doc, ["Role", "Primary surface", "Responsibilities"], [
        ["Platform administrator", "Web admin", "Platform-level oversight, organisation suspension/restoration, user visibility, payment investigation, transfer recovery, health and audit review."],
        ["Organisation owner / integrator", "Web workspace", "Own the subscription and organisation; maintain profile, projects, users, schedules, invoices, payouts and team communication."],
        ["Manager", "Web workspace", "Operate projects, team records, booking and invoice workflows within the same organisation boundary."],
        ["Engineer", "Mobile app", "Maintain field profile, rates and documents; respond to bookings; communicate; start and complete work; create and manage invoices/quotes."],
        ["Guest", "Restricted web", "Access a limited dashboard and permitted conversations without organisation-management powers."],
    ], [1800, 1900, 5660])
    add_callout(doc, "Administrative isolation", "Administrator accounts are control-plane identities and are completely separate from the customer lifecycle. They are not created by public registration, do not belong to a customer organisation, do not start a trial or subscription, and do not consume a plan member allowance.", "info")
    doc.add_heading("Organisation boundary", level=2)
    add_body(doc, "Customer records are scoped to an organisation identifier. Project, user, booking, invoice and document services use that boundary so normal users operate only within their organisation. Platform admin access is separately protected through role-aware middleware, server layouts and API checks; hiding navigation is not treated as authorisation.")

    add_chapter(doc, 3, "Public website and commercial onboarding", "The public experience establishes trust, explains the product and begins the paid customer journey.")
    add_bullets(doc, [
        "Professional homepage positioning Snatchi as an AV project, engineer and booking-management platform.",
        "Dedicated Features, Pricing, About, Contact, Privacy Policy and Terms routes.",
        "Brand ownership presented as Snatchi, owned and operated by PlasmaPro Ltd, with Suftnet as development partner.",
        "Contact workflow using a server endpoint and transactional email integration.",
        "Sign-in, email verification, forgotten-password and password-reset journeys.",
        "Three visible subscription options connected to environment-appropriate Stripe price identifiers.",
        "Stripe checkout success handling and customer billing-portal access for ongoing subscription management.",
        "A public, unindexed tester guide covering web, mobile, plan-boundary and admin acceptance testing."
    ])
    doc.add_heading("Customer onboarding sequence", level=2)
    add_numbered(doc, [
        "A prospective organisation reviews the product and chooses one of the three plans.",
        "The onboarding service validates the selected environment-specific price identifier and prepares the organisation owner account.",
        "Stripe creates and confirms the subscription; webhook handling synchronises the customer and subscription state.",
        "The owner verifies their account, signs in and sees the organisation onboarding checklist.",
        "The checklist guides profile completion, first engineer, first project, first booking and optional payout setup."
    ])

    add_chapter(doc, 4, "Organisation web workspace", "The operational centre for owners and managers.")
    doc.add_heading("Dashboard and onboarding", level=2)
    add_bullets(doc, [
        "Workspace dashboard with recent activity and organisation context.",
        "Progressive onboarding checklist driven by live checks rather than static completion flags.",
        "Direct routes into profile, people, project, booking and payout setup.",
        "Contextual Help & Guides covering the major organisation workflows."
    ])
    doc.add_heading("Projects and documentation", level=2)
    add_bullets(doc, [
        "Create and edit projects with project number, scope, stakeholder, contact, address, priority, status and delivery dates.",
        "Validation for required data, field lengths and valid start/end date ordering.",
        "Attach engineers to project teams and preserve project/team history.",
        "Upload, view and remove project documents through Cloudinary-backed storage.",
        "Search, status visibility and project actions that lead directly into engineer booking."
    ])
    doc.add_heading("People and collaboration", level=2)
    add_bullets(doc, [
        "Create and maintain manager, engineer and guest records with role, active state, visibility and chat access.",
        "Engineer profiles include contact details, addresses, profile imagery and supporting documents.",
        "Direct chat, group chat and booking-linked conversations through Firebase-backed communication services.",
        "Unread-count handling and notification-centre support for operational events."
    ])
    doc.add_heading("Scheduling and engineer sourcing", level=2)
    add_bullets(doc, [
        "Book internal engineers or search the external engineer network by identity and location.",
        "Review engineer availability and service-rate information before making an offer.",
        "Create single-day or linked multi-day bookings while preventing engineer overlaps.",
        "List, filter, inspect and manage schedules with a controlled lifecycle from Pending to terminal outcome.",
        "Approve eligible external bookings and progress them into marketplace payment."
    ])
    doc.add_heading("Invoices, subscriptions and payouts", level=2)
    add_bullets(doc, [
        "Invoice register with search, date filtering, detail review and outcome/status management.",
        "Individual PDF/CSV download and filtered register export.",
        "Organisation profile and password management.",
        "Subscription overview and secure Stripe customer-portal hand-off.",
        "Stripe Connect onboarding and status checks for organisations receiving external-booking payouts.",
        "Views of payments made and payments received, including payment detail."
    ])

    add_chapter(doc, 5, "Engineer mobile application", "A React Native field application for iOS and Android.")
    add_body(doc, "The mobile application is implemented as a native React Native project with iOS and Android targets. It uses the shared Snatchi backend for operational records and Firebase for identity, messaging and push-notification services. The experience is organised around Home, Bookings, Messaging, Invoices and Settings.")
    add_table(doc, ["Area", "Implemented mobile capability"], [
        ["Authentication", "Secure sign-in, persisted authenticated session, verification/forgotten-password support, lock screen and sign-out."],
        ["Home", "Engineer overview, profile access, operational shortcuts and notification entry points."],
        ["Bookings", "Grouped and filtered booking lists, status counts, detail view, Accept/Decline, Start Job and Complete Job actions."],
        ["Project detail", "Project scope, booking context, site information, status progress and project document access."],
        ["Messaging", "Conversation list and real-time chat for direct, group and booking-linked communication."],
        ["Invoices & quotes", "Create, edit and inspect invoice/quote records; status-aware actions; share and download generated PDF documents."],
        ["Profile & compliance", "Edit profile, maintain address, upload/manage certificates and documents, and review images."],
        ["Service rates", "Create, edit and remove hourly, daily or job-based rates used in external sourcing and financial workflows."],
        ["Notifications", "FCM device registration, push receipt, in-app notification list, unread booking handling and deep-link context."],
        ["Location-aware delivery", "Permission-controlled site geofencing starts when eligible work moves In Progress and stops on completion."],
        ["Help", "Mobile help centre for common engineer workflows."],
    ], [2200, 7160], 8.7)
    add_callout(doc, "Mobile release note", "Implemented source capability is not the same as App Store or Play Store release certification. Device-matrix testing, production signing, store metadata/privacy declarations, permission review, crash monitoring and release-build acceptance remain required for launch.", "caution")

    add_chapter(doc, 6, "Connected end-to-end workflow", "How web, mobile and payment services cooperate around one AV job.")
    add_numbered(doc, [
        "Project setup — the organisation creates a project with scope, customer, site, dates, priority and supporting documents.",
        "Engineer selection — an owner or manager chooses an internal engineer or searches the external network, checks availability and reviews rates.",
        "Booking offer — Snatchi creates a schedule record and notifies the engineer through in-app and push channels.",
        "Engineer response — the engineer accepts or declines in the mobile app; the organisation receives the corresponding update.",
        "Approval and payment — eligible external work is approved by the booking organisation and progresses through Stripe payment before it is ready to start.",
        "Field execution — the engineer starts the job in the app; permission-controlled geofencing can activate for the booking site.",
        "Completion — the engineer completes the work, the status is synchronised, and tracking stops.",
        "Financial record — invoice, payment and transfer records remain linked to the engineer, project and schedule for operational review."
    ])
    doc.add_heading("Booking status model", level=2)
    add_table(doc, ["Stage", "Meaning / next responsible party"], [
        ["Pending", "Offer created; engineer must respond."],
        ["Accepted", "Engineer agreed; booking organisation performs the next required action."],
        ["Declined", "Engineer refused; terminal history is preserved."],
        ["Approved / Awaiting Payment", "External booking approved and payment is due."],
        ["Paid / Ready to Start", "Financial prerequisites have completed; engineer can begin at the valid time."],
        ["In Progress", "Engineer has started field work; location-aware workflow may be active."],
        ["Completed", "Engineer confirmed completion; operational workflow is terminal."],
        ["Cancelled / Payment Failed", "Terminal exception path retained for audit and support."],
    ], [2400, 6960])

    add_chapter(doc, 7, "Commercial model and enforceable plans", "A simple launch offer with measurable usage boundaries.")
    add_table(doc, ["Plan", "Price / cycle", "Active projects", "Active members", "Uploads / period", "Max file"], [
        ["Basic Plan", "£50 monthly", "25", "10", "100", "10MB"],
        ["Premium", "£250 / 6 months", "75", "30", "500", "25MB"],
        ["Premium Plus", "£500 yearly", "250", "100", "2,000", "50MB"],
    ], [1600, 1800, 1450, 1450, 1650, 1410], 8.5)
    add_body(doc, "Active-member usage includes the organisation owner, managers and engineers. Platform administrators are excluded. Active projects and active members are calculated from organisation data; document uploads are tracked within the billing period. Per-file size is checked before storage.")
    doc.add_heading("How enforcement works", level=2)
    add_bullets(doc, [
        "The backend resolves the plan from the organisation’s recorded plan or Stripe price identifier.",
        "Protected lifecycle operations require a recognised plan and an active subscription state.",
        "Project, member and upload creation checks compare current usage to the plan limit before committing the operation.",
        "Document size is validated independently of the period upload count.",
        "Usage and entitlement information can be returned for transparent UI messaging and support investigation.",
        "Live Stripe price IDs remain the existing actual Stripe identifiers; test and production IDs are selected by environment."
    ])

    add_chapter(doc, 8, "Payments, notifications and shared services", "The integration layer connecting commercial and operational events.")
    doc.add_heading("Stripe", level=2)
    add_bullets(doc, [
        "Subscription creation, upgrades, cancellation and billing-portal access.",
        "Webhook deduplication and audit support for subscription and marketplace events.",
        "Payment Intents for eligible external engineer bookings.",
        "Stripe Connect onboarding, capability/status checks and destination transfers.",
        "Payment reconciliation records, failure investigation and guarded administrative retry."
    ])
    doc.add_heading("Notifications and email", level=2)
    add_bullets(doc, [
        "Operational events for booking created, accepted, declined, approved, paid, failed, ready, started, completed, updated and cancelled.",
        "Role-aware in-app notifications with related project, schedule, payment and organisation context.",
        "Firebase Cloud Messaging for mobile push and registered device tokens.",
        "Transactional email templates for verification, subscription welcome/update/cancellation, trial milestones, payment outcomes and contact requests.",
        "Duplicate-event checks around high-risk payment notifications."
    ])
    doc.add_heading("Core platform services", level=2)
    add_table(doc, ["Service", "Purpose"], [
        ["MongoDB / Mongoose", "Persistent organisations, users, projects, bookings, invoices, payments, notifications, entitlements and audit records."],
        ["Firebase", "Mobile authentication provisioning, direct/group/booking chat and push messaging."],
        ["Cloudinary", "Project and engineer document/image storage with metadata and cleanup support."],
        ["Brevo / email", "Transactional customer, subscription and support communication."],
        ["Sentry", "Web and mobile error-monitoring integration."],
    ], [2200, 7160])

    add_chapter(doc, 9, "Platform administration and operational safety", "A premium first release focused on reversible, auditable controls.")
    add_bullets(doc, [
        "Global paginated visibility of organisations and users.",
        "Organisation detail and reversible suspension/restoration without cancelling the Stripe subscription.",
        "Payment-failure register and investigation views with Stripe and reconciliation references.",
        "Manual transfer retry guarded by current payment checks, destination capability checks, deterministic idempotency and a concurrency lock.",
        "Mandatory operator reasons for organisation access changes and transfer retries.",
        "Persistent admin audit records including actor, action, target, reason, result, time, IP address and user agent.",
        "Read-only integration health reporting that indicates configuration state without returning secret values.",
        "Customer-user access is read-only for platform admins in this release; permanent deletion and unaudited profile edits are intentionally excluded."
    ])
    add_callout(doc, "Security milestone not represented as complete", "Admin MFA and advanced session management remain a separate release milestone. This limitation is explicitly documented and should be closed before broad production administration or high-risk operational scale.", "risk")
    doc.add_heading("Administrator provisioning", level=2)
    add_body(doc, "Administrators are provisioned only by a dedicated script with explicit test/live targets, dry-run default, hidden password input, password-strength limits and an additional confirmation for live writes. The script creates or updates only the specified identity and removes any legacy organisation association.")

    add_chapter(doc, 10, "Production-readiness work completed", "Changes made to move the product from functioning software toward a supportable launch.")
    add_bullets(doc, [
        "Repositioned public content and identity around Snatchi as the product, PlasmaPro Ltd as owner/operator and Suftnet as developer.",
        "Simplified the commercial offer into three realistic plans and added server-side entitlement enforcement.",
        "Preserved verified Stripe live price IDs while separating environment-specific price validation.",
        "Repaired and hardened scheduler creation/status paths and expanded booking notifications.",
        "Added organisation onboarding checks, contextual web help and a comprehensive acceptance-testing guide.",
        "Expanded invoice workflows, including web filtering/export and mobile invoice/quote capability.",
        "Hardened authentication, admin access control, webhook processing and payment reconciliation/recovery paths.",
        "Built a safer admin dashboard with reversible access controls, health visibility and audit records.",
        "Upgraded the web framework and direct dependencies, removed known critical Yarn audit findings and documented residual transitive risk.",
        "Validated the web application with a successful production build and focused admin access-control tests during the readiness pass."
    ])
    doc.add_heading("Verification position", level=2)
    add_table(doc, ["Area", "Current evidence", "Position"], [
        ["Web production compilation", "Next.js 15.5.23 production build completed successfully during the latest guide implementation and prior readiness pass.", "Passed"],
        ["Admin access control", "Focused unauthenticated/non-admin/admin scenarios passed in the production-readiness pass.", "Passed"],
        ["Critical dependency advisories", "Known critical findings reported by Yarn Classic were removed; residual high transitive tooling findings are documented.", "Managed / monitor"],
        ["Mobile release", "iOS/Android source, native projects and core workflows are present; store-release certification is not evidenced by this report.", "Final QA required"],
        ["Live services", "Actual production environment, database, Stripe, Firebase, Cloudinary, email and monitoring configuration must be verified at deployment.", "Deployment gate"],
    ], [1900, 5600, 1860], 8.5)

    add_chapter(doc, 11, "Recommended path to launch", "A controlled sequence for converting the implemented product into a dependable live service.")
    add_numbered(doc, [
        "Freeze the release candidate — confirm web and mobile commit/version, environments and release owners.",
        "Rotate and validate production credentials — MongoDB rotation is already planned; verify all other secrets and least-privilege access.",
        "Complete the acceptance guide — run public web, owner/manager, engineer mobile, cross-role booking, all plan boundaries and authorised admin scenarios.",
        "Validate Stripe in both modes — subscriptions, customer portal, Connect onboarding, external booking payment, webhook idempotency and payout recovery.",
        "Certify mobile release builds — physical iOS/Android devices, permissions, push notifications, geofencing, deep links, crash reporting and store declarations.",
        "Complete admin security milestone — MFA, session invalidation/expiry controls and operator recovery procedures.",
        "Perform production-safe smoke tests — test records only, audit every privileged action and verify monitoring/alerts.",
        "Launch with support ownership — named incident lead, customer support route, backup/restore confidence, payment reconciliation routine and dependency-review cadence."
    ])
    doc.add_heading("Go-live decision criteria", level=2)
    add_bullets(doc, [
        "No open blocker or critical defect in account access, project creation, booking lifecycle, payment or data isolation.",
        "All three plans display and enforce the same limits in test and production configuration.",
        "Mobile release builds complete the coordinated booking workflow on supported devices.",
        "Production administrators use isolated identities and the agreed MFA/session controls.",
        "Stripe webhooks, email, push notifications, document storage and monitoring are verified with production-safe test records.",
        "PlasmaPro has documented ownership for operations, billing support, security incidents and customer communication."
    ])

    add_chapter(doc, 12, "Conclusion", "The platform foundation is coherent, commercially understandable and ready for disciplined final release work.")
    add_body(doc, "Snatchi now represents more than a collection of project screens. It is a connected operating model for AV service delivery: office teams organise work on the web, engineers execute it in the field through mobile, and shared booking, payment, notification and audit services keep the lifecycle coherent.")
    add_body(doc, "The strongest launch story for PlasmaPro is practical: Snatchi reduces fragmentation between project planning, engineer coordination and job completion while providing a clear subscription model and a safer operational control plane. The remaining work is principally release assurance and security completion, not reinvention of the core product.")
    add_callout(doc, "Recommended presentation message", "Snatchi gives AV organisations one connected workspace for projects, people and bookings, while giving engineers a purpose-built mobile tool for field delivery. PlasmaPro owns a credible platform foundation that can launch with a focused core and expand incrementally from real customer usage.", "info")
    doc.add_heading("Appendix A — principal routes and product areas", level=2)
    add_table(doc, ["Product area", "Representative location"], [
        ["Public website", "/, /features, /pricing, /about, /contact"],
        ["Organisation dashboard", "/protected/integrator/dashboard"],
        ["Projects and team", "/protected/integrator/project, /protected/integrator/user"],
        ["Scheduling and sourcing", "/protected/integrator/scheduler, /search-engineers"],
        ["Chat and invoices", "/protected/integrator/chat, /protected/integrator/invoice"],
        ["Settings and help", "/protected/integrator/settings, /protected/integrator/help"],
        ["Admin control plane", "/protected/admin/dashboard, /integrator, /user, /payments/failures"],
        ["Acceptance guide", "/testing-guide"],
        ["Mobile tabs", "Home, Bookings, Messaging, Invoices, Settings"],
    ], [2900, 6460], 8.7)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    r = p.add_run("Document basis")
    set_font(r, 9, True, INDIGO)
    p = doc.add_paragraph()
    r = p.add_run("Prepared from the Snatchi web and React Native source repositories, current application routes, service models, product guide, production-readiness README and recent implementation history as inspected on 22 August 2026. This is a product and delivery overview, not an independent security certification or financial assurance report.")
    set_font(r, 8.5, False, MUTED)

    # Keep rows intact when possible, without imposing fixed heights.
    for table in doc.tables:
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            cant_split = OxmlElement("w:cantSplit")
            tr_pr.append(cant_split)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "Snatchi Product Overview and Delivery Position"
    doc.core_properties.subject = "PlasmaPro stakeholder presentation"
    doc.core_properties.author = "Suftnet"
    doc.core_properties.keywords = "Snatchi, PlasmaPro, AV, project management, engineer booking, mobile"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
