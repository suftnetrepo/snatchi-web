from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement

from generateProductOverview import (
    INDIGO, DEEP, INK, MUTED,
    set_font, page_number, border_bottom, add_field, add_body, add_bullets,
    add_numbered, add_callout, add_table
)

OUT = Path(__file__).resolve().parents[1] / "docs" / "Snatchi-Plan-Pricing-Recommendation-for-PlasmaPro.docx"


def add_chapter(doc, number, title, deck=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"{number:02d}  /  PRICING RECOMMENDATION")
    set_font(r, 8.5, True, INDIGO)
    doc.add_heading(title, level=1)
    if deck:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run(deck)
        set_font(r, 10.5, False, MUTED)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(.82)
    section.bottom_margin = Inches(.82)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(.42)
    section.footer_distance = Inches(.42)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, before, after, color in (
        ("Heading 1", 16, 16, 8, DEEP),
        ("Heading 2", 13, 12, 6, INDIGO),
        ("Heading 3", 11.5, 8, 4, DEEP),
    ):
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    r = header.add_run("SNATCHI  /  COMMERCIAL PLAN PRICING")
    set_font(r, 8, True, MUTED)
    page_number(section.footer.paragraphs[0])

    # Decision memo masthead.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("COMMERCIAL PRICING MEMO")
    set_font(r, 10, True, INDIGO)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("Recommended Subscription Pricing\nfor Snatchi")
    set_font(r, 28, True, DEEP, "Aptos Display")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("A launch pricing structure designed to cover external operating costs and create room for sustainable growth")
    set_font(r, 13, False, MUTED)
    border_bottom(p, INDIGO, "18")
    add_field(doc, "PREPARED FOR", "PlasmaPro Ltd")
    add_field(doc, "PRODUCT", "Snatchi")
    add_field(doc, "DATE", "22 August 2026")
    add_field(doc, "PRICING BASIS", "External-service cost plan, current plan allowances and Stripe UK fee assumptions")
    add_field(doc, "PRICE PRESENTATION", "Recommended prices exclude VAT")
    add_callout(doc, "Recommendation", "Launch Basic at £79 monthly, Premium at £799 every six months, and Premium Plus at £2,990 yearly. This corrects the current inverse pricing structure, creates visible value separation and allows a small customer base to cover the recommended £190 monthly external-service budget.", "info")

    add_chapter(doc, 1, "Executive recommendation", "Three commercially distinct plans with progressively stronger contribution.")
    add_table(doc, ["Plan", "Recommended charge", "Monthly equivalent", "Net after Stripe*", "Target customer"], [
        ["Basic", "£79 monthly", "£79.00", "£77.06/month", "Small AV teams and owner-led businesses"],
        ["Premium", "£799 every 6 months", "£133.17", "£130.20/month", "Established integrators with several engineers/managers"],
        ["Premium Plus", "£2,990 yearly", "£249.17", "£243.67/month", "Larger organisations and multi-team operations"],
    ], [1300, 1850, 1600, 1800, 2810], 8.5)
    p = doc.add_paragraph()
    r = p.add_run("* Assumes a standard UK card and combines Stripe Payments at 1.5% + £0.20 with Stripe Billing at 0.7%. Excludes VAT, refunds, disputes and labour.")
    set_font(r, 8.3, False, MUTED)
    add_callout(doc, "Why change is necessary", "The current Premium (£250 per six months) and Premium Plus (£500 per year) both cost about £41.67 per month—less than the £50 monthly Basic plan—despite offering substantially more projects, members, uploads and file capacity. That structure weakens upgrade logic and leaves insufficient room for support and growth.", "caution")

    add_chapter(doc, 2, "Recommended plan structure", "The customer should understand the difference without needing a sales explanation.")
    doc.add_heading("Basic — £79 per month", level=2)
    add_bullets(doc, [
        "Up to 25 active projects.",
        "Up to 10 active organisation members, including the owner, managers and engineers.",
        "Up to 100 document uploads per billing period and 10MB per file.",
        "All core project, booking, communication and invoice workflows.",
        "Standard email support during business hours.",
        "Positioning: a complete operational starting point, not a restricted trial product."
    ])
    doc.add_heading("Premium — £799 every six months", level=2)
    add_bullets(doc, [
        "Up to 75 active projects.",
        "Up to 30 active organisation members.",
        "Up to 500 document uploads per billing period and 25MB per file.",
        "All core Snatchi functionality.",
        "Priority email support.",
        "Positioning: the recommended plan for established AV integrators running several concurrent teams."
    ])
    doc.add_heading("Premium Plus — £2,990 per year", level=2)
    add_bullets(doc, [
        "Up to 250 active projects.",
        "Up to 100 active organisation members.",
        "Up to 2,000 document uploads per billing period and 50MB per file.",
        "All core Snatchi functionality.",
        "Priority support and the natural home for future advanced capabilities.",
        "Positioning: a high-capacity operational platform for larger organisations, not a discounted version of Premium."
    ])
    add_body(doc, "The commercial difference is based on measurable capacity rather than speculative features. This keeps launch promises enforceable while allowing PlasmaPro to add higher-value capabilities incrementally.")

    add_chapter(doc, 3, "Contribution and external-cost coverage", "What each customer contributes after estimated Stripe subscription fees.")
    add_table(doc, ["Plan", "Customer charge", "Est. Stripe fee", "Net per billing event", "Net monthly equivalent"], [
        ["Basic", "£79", "£1.94", "£77.06", "£77.06"],
        ["Premium", "£799", "£17.78", "£781.22", "£130.20"],
        ["Premium Plus", "£2,990", "£65.98", "£2,924.02", "£243.67"],
    ], [1400, 1750, 1650, 2200, 2360], 8.5)
    add_body(doc, "Against the recommended £190 monthly launch external-service budget, the approximate standalone break-even point is three Basic customers, two Premium customers, or one Premium Plus customer. This is technical-service break-even only; a viable business must also cover labour, customer support, tax, sales and risk.")
    doc.add_heading("Illustrative customer-mix scenarios", level=2)
    add_table(doc, ["Scenario", "Customer mix", "Net subscription revenue", "External services", "Contribution before labour"], [
        ["Minimum technical coverage", "3 Basic", "~£231/month", "£190/month", "~£41/month"],
        ["Balanced launch", "10 Basic + 3 Premium + 1 Premium Plus", "~£1,405/month", "£190/month", "~£1,215/month"],
        ["Growth base", "20 Basic + 8 Premium + 2 Premium Plus", "~£3,070/month", "£420/month", "~£2,650/month"],
    ], [1600, 2850, 1900, 1500, 1510], 8.2)
    add_callout(doc, "Interpretation", "Contribution before labour is not net profit. The balanced launch scenario leaves meaningful room for support and product work, but PlasmaPro must add its real monthly labour, sales, tax and overhead costs before setting a final profit target.", "risk")

    add_chapter(doc, 4, "Founding-customer launch offer", "Use a time-limited incentive without permanently damaging the standard price.")
    add_table(doc, ["Plan", "Founding offer", "Standard renewal", "Maximum offer period"], [
        ["Basic", "£59/month", "£79/month", "First 12 months"],
        ["Premium", "£599/6 months", "£799/6 months", "First 6-month term"],
        ["Premium Plus", "£2,390/year", "£2,990/year", "First annual term"],
    ], [1500, 2100, 2100, 3660], 8.7)
    add_bullets(doc, [
        "Show the standard renewal price before purchase and in the order confirmation.",
        "Limit eligibility by customer count or a clear campaign closing date.",
        "Do not offer lifetime pricing or permanent percentage discounts.",
        "Do not combine the founding offer with custom free months unless approved individually.",
        "Measure activation, retention, support demand and plan-limit usage before extending the campaign."
    ])
    add_body(doc, "Estimated founding-offer net monthly equivalents after the same Stripe assumption are approximately £57.50 for Basic, £97.60 for Premium and £194.77 for Premium Plus. The offer therefore remains above direct technical-service cost at a modest customer count while reducing purchase friction.")

    add_chapter(doc, 5, "Pricing governance", "Rules for protecting margin after launch.")
    add_numbered(doc, [
        "Publish all prices excluding VAT where appropriate and state the billing term prominently.",
        "Keep Stripe transaction fees and external-engineer marketplace fees separate in management reporting.",
        "Review acquisition, activation, retention, support time and infrastructure cost at 30, 60 and 90 days.",
        "Measure gross margin by plan; do not average a costly large customer into the whole customer base.",
        "Use entitlement usage to identify when limits are too generous, too restrictive or poorly understood.",
        "Review pricing annually and provide fair notice before changing renewal prices.",
        "Reserve custom enterprise pricing for requirements that materially change support, security, storage or integration cost."
    ])
    doc.add_heading("Target economics", level=2)
    add_bullets(doc, [
        "External services should remain a small proportion of subscription revenue once the customer base is established.",
        "PlasmaPro should model labour and support before adopting a final gross-margin target; a 75–80% software gross-margin objective is a useful long-term management target, not a guarantee from this report.",
        "Payment failures, refunds, disputes and bad debt require a separate reserve.",
        "Annual and six-month plans improve cash collection but create a service obligation for the full prepaid period.",
        "Premium Plus should not be discounted merely because it bills annually; its capacity can create materially higher storage, support and operational exposure."
    ])

    add_chapter(doc, 6, "Approval recommendation", "A clear decision for PlasmaPro.")
    add_callout(doc, "Approve for launch", "Basic: £79 monthly. Premium: £799 every six months. Premium Plus: £2,990 yearly. Prices exclude VAT. Permit the documented founding-customer offer for one initial term only, and review pricing after 90 days of real customer behaviour.", "positive")
    add_body(doc, "This structure gives Snatchi a credible B2B SaaS price ladder while preserving a relatively accessible entry plan. It aligns price with project, member and document capacity, supports the recommended external-service budget, and leaves room for PlasmaPro to fund the human work that is outside the operating-cost report.")
    add_body(doc, "The recommendation is a commercial planning model, not a guarantee of profitability. Final approval should combine this report with PlasmaPro’s labour cost, expected customer acquisition cost, tax treatment, support commitment and revenue forecast.")

    doc.add_heading("Appendix — calculation assumptions", level=2)
    add_table(doc, ["Assumption", "Value"], [
        ["Launch external-service budget", "£190/month excluding VAT and Stripe fees"],
        ["Growth external-service budget", "£420/month excluding VAT and Stripe fees"],
        ["Standard UK Stripe card fee", "1.5% + £0.20 per successful payment"],
        ["Stripe Billing fee", "0.7% of recurring Billing volume"],
        ["Combined planning formula", "2.2% of charge + £0.20"],
        ["VAT", "Excluded from all recommended prices and calculations"],
        ["Manual labour and overhead", "Excluded; must be added by PlasmaPro"],
        ["Plan allowances", "Current enforced project, member, upload and file-size limits"],
    ], [3000, 6360], 8.7)

    for table in doc.tables:
        for row in table.rows:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "Snatchi Subscription Plan Pricing Recommendation"
    doc.core_properties.subject = "Commercial pricing recommendation for PlasmaPro"
    doc.core_properties.author = "Suftnet"
    doc.core_properties.keywords = "Snatchi, PlasmaPro, pricing, subscription, profitability, SaaS"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
